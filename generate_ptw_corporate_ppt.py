import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_ptw_presentation(output_paths):
    prs = Presentation()
    # 16:9 Widescreen: 13.333 x 7.5 inches
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Enterprise Palette: Clean, professional, industrial safety aesthetic
    NAVY = RGBColor(15, 30, 54)          # Primary dark typography & solid emphasis
    DEEP_BLUE = RGBColor(27, 54, 93)     # Secondary headers
    ACCENT_TEAL = RGBColor(14, 116, 144) # Single restrained accent (Slate Cyan/Teal)
    SLATE_GRAY = RGBColor(71, 85, 105)   # Subtitles & body text
    LIGHT_BG = RGBColor(248, 250, 252)   # Light neutral background
    CARD_BG = RGBColor(255, 255, 255)    # Card fill
    CARD_BORDER = RGBColor(203, 213, 225)# Subtle card border
    HIGHLIGHT_BG = RGBColor(241, 245, 249) # Neutral container fill
    WHITE = RGBColor(255, 255, 255)
    AMBER_DEFER = RGBColor(180, 83, 9)   # Controlled amber for Deferral flow
    BORDER_CYAN = RGBColor(14, 116, 144) # Card accent border

    def setup_slide_header(slide, slide_num, title_text, takeaway_subtext):
        # Base background
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = LIGHT_BG
        bg.line.fill.background()

        # Slim accent line at the top
        top_accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.04))
        top_accent.fill.solid()
        top_accent.fill.fore_color.rgb = ACCENT_TEAL
        top_accent.line.fill.background()

        # Header Text Box
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.55), Inches(11.733), Inches(1.15))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = title_text
        p1.font.size = Pt(32)
        p1.font.bold = True
        p1.font.color.rgb = NAVY
        p1.font.name = "Calibri"

        if takeaway_subtext:
            p2 = tf.add_paragraph()
            p2.text = takeaway_subtext
            p2.font.size = Pt(14)
            p2.font.color.rgb = SLATE_GRAY
            p2.font.name = "Calibri"
            p2.space_before = Pt(3)

    def add_bottom_takeaway(slide, text):
        tb = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.6), Inches(11.733), Inches(0.55))
        tb.fill.solid()
        tb.fill.fore_color.rgb = HIGHLIGHT_BG
        tb.line.color.rgb = CARD_BORDER
        tb.line.width = Pt(1)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(13.5)
        p.font.bold = True
        p.font.color.rgb = NAVY
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Calibri"

    # =========================================================================
    # SLIDE 1: 1. What is the Application About?
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    setup_slide_header(
        s1, "1", "1. What is the Application About?",
        "Key Message: A centralized platform for controlling and tracking work activities from permit creation through execution and closure."
    )

    # Core Value Pillars Banner (Work Authorization + Safety Control + Operational Visibility + Accountability + Traceability)
    pillars = ["WORK AUTHORIZATION", "+", "SAFETY CONTROL", "+", "OPERATIONAL VISIBILITY", "+", "ACCOUNTABILITY", "+", "TRACEABILITY"]
    p_box = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.85), Inches(11.733), Inches(0.52))
    p_box.fill.solid()
    p_box.fill.fore_color.rgb = NAVY
    p_box.line.fill.background()
    p_tf = p_box.text_frame
    p_tf.word_wrap = True
    p_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_para = p_tf.paragraphs[0]
    p_para.text = "   ".join(pillars)
    p_para.font.size = Pt(13)
    p_para.font.bold = True
    p_para.font.color.rgb = WHITE
    p_para.alignment = PP_ALIGN.CENTER
    p_para.font.name = "Calibri"

    # Central Architecture Visual: PTW Platform -> [AUTHORIZE | CONTROL | MONITOR]
    col_w = Inches(3.75)
    col_gap = Inches(0.24)
    cols_data = [
        ("AUTHORIZE", [("Approval", "Multi-stage queue"), ("Execution", "Active permit control"), ("Evidence", "Verification uploads"), ("Closure", "Site sign-off"), ("Analytics", "KPI reporting")]),
        ("CONTROL", [("LOTOTO Plans", "Energy isolation & lockouts"), ("Safety Controls", "Pre-task safety checklists"), ("Active Oversight", "Daily status monitoring"), ("Permit Rules", "Mandatory work conditions"), ("Safe Closure", "Handback verification")]),
        ("MONITOR", [("SIMOPS", "Real-time conflict detection"), ("Incidents", "Near miss & unsafe tracking"), ("Operational KPIs", "Live status dashboard"), ("Audit Trail", "Full history log"), ("Reports", "Operational data exports")])
    ]

    for idx, (head, items) in enumerate(cols_data):
        c_left = Inches(0.8) + idx * (col_w + col_gap)
        card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, Inches(2.52), col_w, Inches(3.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1.5)

        # Column Header Strip
        ch = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, c_left, Inches(2.52), col_w, Inches(0.45))
        ch.fill.solid()
        ch.fill.fore_color.rgb = ACCENT_TEAL if idx == 1 else DEEP_BLUE
        ch.line.fill.background()
        ch_tf = ch.text_frame
        ch_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        ch_p = ch_tf.paragraphs[0]
        ch_p.text = head
        ch_p.font.size = Pt(14)
        ch_p.font.bold = True
        ch_p.font.color.rgb = WHITE
        ch_p.alignment = PP_ALIGN.CENTER
        ch_p.font.name = "Calibri"

        # Column Content
        tb = s1.shapes.add_textbox(c_left + Inches(0.25), Inches(3.05), col_w - Inches(0.5), Inches(2.5))
        ctf = tb.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = ctf.margin_top = ctf.margin_bottom = 0
        for item_idx, (t, d) in enumerate(items):
            ip = ctf.add_paragraph() if item_idx > 0 else ctf.paragraphs[0]
            ip.text = f"• {t}  —  {d}"
            ip.font.size = Pt(12)
            ip.font.bold = (item_idx == 0)
            ip.font.color.rgb = NAVY if item_idx == 0 else SLATE_GRAY
            ip.font.name = "Calibri"
            ip.space_after = Pt(6)

    # Key Roles Bar
    roles_box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.85), Inches(11.733), Inches(0.6))
    roles_box.fill.solid()
    roles_box.fill.fore_color.rgb = WHITE
    roles_box.line.color.rgb = CARD_BORDER
    roles_box.line.width = Pt(1)
    rtf = roles_box.text_frame
    rtf.word_wrap = True
    rtf.vertical_anchor = MSO_ANCHOR.MIDDLE
    rp = rtf.paragraphs[0]
    rp.text = "Accountable Platform Roles:   Platform Admin   |   Org Admin   |   Job Issuer   |   Safety Officer   |   Operator   |   Viewer"
    rp.font.size = Pt(13)
    rp.font.bold = True
    rp.font.color.rgb = NAVY
    rp.alignment = PP_ALIGN.CENTER
    rp.font.name = "Calibri"

    add_bottom_takeaway(s1, "Takeaway: PTW acts as the central operational record connecting work authorization, safety controls, execution, and closure.")

    # =========================================================================
    # SLIDE 2: 2. Features of the Application
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    setup_slide_header(
        s2, "2", "2. Features of the Application",
        "Key Message: Built-in capabilities grouped across the core business and operational functions of the enterprise."
    )

    feature_groups = [
        ("1. WORK AUTHORIZATION", ACCENT_TEAL, [
            "Permit creation & templates",
            "Draft lifecycle management",
            "Structured permit submission",
            "Multi-stage approval queue",
            "Approve or Defer with feedback",
            "Resubmission of deferred permits"
        ]),
        ("2. SAFETY & OPERATIONAL CONTROL", DEEP_BLUE, [
            "LOTOTO plans & energy isolation",
            "SIMOPS live conflict monitoring",
            "Safety-related controls & checklists",
            "Workplace incident reporting",
            "Near-miss reporting & tracking",
            "Unsafe-condition logging"
        ]),
        ("3. WORK EXECUTION & TRACEABILITY", ACCENT_TEAL, [
            "Approved work execution state",
            "Live progress updates & notes",
            "Execution evidence upload",
            "Auditable activity history log",
            "Suspend work capability",
            "Structured Verify & Close protocol",
            "Historical permits archive"
        ]),
        ("4. MANAGEMENT & VISIBILITY", DEEP_BLUE, [
            "Workforce & user management",
            "Roles, permissions & competencies",
            "Organization hierarchy & plants",
            "Operational KPIs & live analytics",
            "Reports and CSV/data exports",
            "Billing & subscription tracking"
        ])
    ]

    card_w2 = Inches(2.78)
    card_gap2 = Inches(0.2)
    card_top2 = Inches(1.85)
    card_h2 = Inches(4.55)

    for i, (grp_title, header_col, items) in enumerate(feature_groups):
        c_left = Inches(0.8) + i * (card_w2 + card_gap2)
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, card_top2, card_w2, card_h2)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1.5)

        # Card top banner
        hdr = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, c_left, card_top2, card_w2, Inches(0.55))
        hdr.fill.solid()
        hdr.fill.fore_color.rgb = header_col
        hdr.line.fill.background()
        htf = hdr.text_frame
        htf.vertical_anchor = MSO_ANCHOR.MIDDLE
        hp = htf.paragraphs[0]
        hp.text = grp_title
        hp.font.size = Pt(12.5)
        hp.font.bold = True
        hp.font.color.rgb = WHITE
        hp.alignment = PP_ALIGN.CENTER
        hp.font.name = "Calibri"

        # Card items
        tb = s2.shapes.add_textbox(c_left + Inches(0.18), card_top2 + Inches(0.68), card_w2 - Inches(0.36), card_h2 - Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        for item_idx, itm in enumerate(items):
            p = tf.add_paragraph() if item_idx > 0 else tf.paragraphs[0]
            p.text = f"✔  {itm}"
            p.font.size = Pt(12)
            p.font.color.rgb = NAVY
            p.font.name = "Calibri"
            p.space_after = Pt(8)

    add_bottom_takeaway(s2, "One Platform  →  From work authorization to operational reporting and compliance")

    # =========================================================================
    # SLIDE 3: 3. Why is this Application Required?
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    setup_slide_header(
        s3, "3", "3. Why is this Application Required?",
        "Key Message: Provides a structured digital process to support safer, accountable, and more controlled work operations."
    )

    # Comparison Columns: Without Control vs. With PTW Application
    # Left Comparison Box: Without Control (Risks)
    comp_w = Inches(3.4)
    box_wo = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.85), comp_w, Inches(4.55))
    box_wo.fill.solid()
    box_wo.fill.fore_color.rgb = HIGHLIGHT_BG
    box_wo.line.color.rgb = CARD_BORDER
    box_wo.line.width = Pt(1.5)

    wo_hdr = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.85), comp_w, Inches(0.5))
    wo_hdr.fill.solid()
    wo_hdr.fill.fore_color.rgb = SLATE_GRAY
    wo_hdr.line.fill.background()
    wo_htf = wo_hdr.text_frame
    wo_htf.vertical_anchor = MSO_ANCHOR.MIDDLE
    wo_hp = wo_htf.paragraphs[0]
    wo_hp.text = "WITHOUT STRUCTURED CONTROL"
    wo_hp.font.size = Pt(13)
    wo_hp.font.bold = True
    wo_hp.font.color.rgb = WHITE
    wo_hp.alignment = PP_ALIGN.CENTER
    wo_hp.font.name = "Calibri"

    wo_tb = s3.shapes.add_textbox(Inches(1.0), Inches(2.5), comp_w - Inches(0.4), Inches(3.7))
    wo_tf = wo_tb.text_frame
    wo_tf.word_wrap = True
    wo_tf.margin_left = wo_tf.margin_right = wo_tf.margin_top = wo_tf.margin_bottom = 0

    wo_items = [
        ("Fragmented Information", "Paper permits, disconnected spreadsheets, and informal communications."),
        ("Limited Operational Visibility", "Management cannot easily monitor concurrent active worksites or statuses."),
        ("Higher Coordination Risk", "Unidentified schedule/spatial overlaps lead to dangerous simultaneous tasks."),
        ("Audit & Accountability Gaps", "Difficult to verify who approved, what controls were verified, or past history.")
    ]
    for idx, (t, d) in enumerate(wo_items):
        p1 = wo_tf.add_paragraph() if idx > 0 else wo_tf.paragraphs[0]
        p1.text = f"✖  {t}"
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = NAVY
        p1.font.name = "Calibri"
        p2 = wo_tf.add_paragraph()
        p2.text = d
        p2.font.size = Pt(11)
        p2.font.color.rgb = SLATE_GRAY
        p2.font.name = "Calibri"
        p2.space_after = Pt(10)

    # Right Comparison Box: 6 Business Outcomes (With PTW Application)
    outcomes_w = Inches(8.0)
    outcomes_left = Inches(4.533)
    
    outcomes = [
        ("1. CONTROLLED WORK", "Structured lifecycle governing preparation, review, authorization, and closure before tools touch equipment."),
        ("2. SAFER OPERATIONS", "Integrated safety controls including mandatory LOTOTO plans, pre-task checklists, and SIMOPS conflict monitoring."),
        ("3. CLEAR ACCOUNTABILITY", "Role-based permissions make explicit who authored, inspected, approved, executed, and verified each work order."),
        ("4. OPERATIONAL VISIBILITY", "Single-pane dashboard into active permits, pending approvals, execution stages, incidents, LOTOTO, and SIMOPS."),
        ("5. TRACEABILITY", "Real-time execution updates, photographic evidence upload, immutable activity history, and searchable closed archives."),
        ("6. DATA-DRIVEN MANAGEMENT", "Live KPI dashboards and operational CSV reports provide data for continuous safety review and operational optimization.")
    ]

    # Render as a 2x3 grid of clean cards
    card_grid_w = Inches(3.88)
    card_grid_h = Inches(1.38)
    grid_gap_x = Inches(0.24)
    grid_gap_y = Inches(0.18)

    for i, (title, desc) in enumerate(outcomes):
        col_idx = i % 2
        row_idx = i // 2
        c_x = outcomes_left + col_idx * (card_grid_w + grid_gap_x)
        c_y = Inches(1.85) + row_idx * (card_grid_h + grid_gap_y)

        c = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x, c_y, card_grid_w, card_grid_h)
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = CARD_BORDER
        c.line.width = Pt(1.5)

        # Small left accent bar
        bar = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, c_x, c_y, Inches(0.12), card_grid_h)
        bar.fill.solid()
        bar.fill.fore_color.rgb = ACCENT_TEAL
        bar.line.fill.background()

        tb = s3.shapes.add_textbox(c_x + Inches(0.24), c_y + Inches(0.12), card_grid_w - Inches(0.35), card_grid_h - Inches(0.24))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(12.5)
        p1.font.bold = True
        p1.font.color.rgb = DEEP_BLUE
        p1.font.name = "Calibri"

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = SLATE_GRAY
        p2.font.name = "Calibri"
        p2.space_before = Pt(2)

    add_bottom_takeaway(s3, "Business Impact: Provides a structured digital process to support safer, auditable, and more controlled work operations.")

    # =========================================================================
    # SLIDE 4: 4. End-to-End Application Workflow
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    setup_slide_header(
        s4, "4", "4. End-to-End Application Workflow",
        "Key Message: The horizontal lifecycle moving work orders safely from initial creation to final verified archive."
    )

    # Status Bar Flow Prominent across top of lifecycle
    status_box = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.85), Inches(11.733), Inches(0.48))
    status_box.fill.solid()
    status_box.fill.fore_color.rgb = NAVY
    status_box.line.fill.background()
    s_tf = status_box.text_frame
    s_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    s_p = s_tf.paragraphs[0]
    s_p.text = "OFFICIAL STATUS FLOW:   DRAFT   →   PENDING APPROVAL   →   APPROVED   →   EXECUTION   →   CLOSED"
    s_p.font.size = Pt(13)
    s_p.font.bold = True
    s_p.font.color.rgb = WHITE
    s_p.alignment = PP_ALIGN.CENTER
    s_p.font.name = "Calibri"

    # Main 8 Horizontal Stages
    # Layout: 8 stages in a structured, readable sequence
    # Top row: 1 CREATE, 2 PLAN, 3 SUBMIT, 4 APPROVE (with side DEFERRED loop)
    # Bottom row: 5 CONTROL, 6 EXECUTE, 7 VERIFY, 8 ARCHIVE
    # Or 8 sequential steps across 2 rows of 4 with clear interconnecting flow
    flow_steps = [
        ("STAGE 1", "CREATE", "Permit Draft", "Job Issuer initiates new permit instance", ACCENT_TEAL),
        ("STAGE 2", "PLAN", "Work Scope & Plan", "Basic info, plant location, schedule & crew", DEEP_BLUE),
        ("STAGE 3", "SUBMIT", "Review & Submit", "Validation check; moves to Pending Approval", ACCENT_TEAL),
        ("STAGE 4", "APPROVE", "Approval Queue", "Authorized role approves OR defers permit", DEEP_BLUE),
        ("STAGE 5", "CONTROL", "LOTOTO & SIMOPS", "Energy isolation checks & conflict review", ACCENT_TEAL),
        ("STAGE 6", "EXECUTE", "Work Execution", "Progress updates, evidence & suspend option", DEEP_BLUE),
        ("STAGE 7", "VERIFY", "Verify & Close", "Work site inspection & structured sign-off", ACCENT_TEAL),
        ("STAGE 8", "ARCHIVE", "Historical Record", "Closed record stored for analytics & audit", DEEP_BLUE),
    ]

    card_w4 = Inches(2.78)
    card_h4 = Inches(1.7)
    row1_y = Inches(2.55)
    row2_y = Inches(4.7)

    for i, (stg_num, stg_title, stg_sub, stg_desc, col) in enumerate(flow_steps):
        row = i // 4
        col_idx = i % 4
        left = Inches(0.8) + col_idx * (card_w4 + card_gap2)
        top = row1_y if row == 0 else row2_y

        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, card_w4, card_h4)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1.5)

        # Stage Number Badge
        badge = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.12), top + Inches(0.12), Inches(0.85), Inches(0.3))
        badge.fill.solid()
        badge.fill.fore_color.rgb = col
        badge.line.fill.background()
        btf = badge.text_frame
        btf.vertical_anchor = MSO_ANCHOR.MIDDLE
        bp = btf.paragraphs[0]
        bp.text = stg_num
        bp.font.size = Pt(10)
        bp.font.bold = True
        bp.font.color.rgb = WHITE
        bp.alignment = PP_ALIGN.CENTER
        bp.font.name = "Calibri"

        # Stage Text
        tb = s4.shapes.add_textbox(left + Inches(0.15), top + Inches(0.48), card_w4 - Inches(0.3), card_h4 - Inches(0.55))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = stg_title
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = NAVY
        p1.font.name = "Calibri"

        p2 = tf.add_paragraph()
        p2.text = stg_sub
        p2.font.size = Pt(11.5)
        p2.font.bold = True
        p2.font.color.rgb = col
        p2.font.name = "Calibri"

        p3 = tf.add_paragraph()
        p3.text = stg_desc
        p3.font.size = Pt(10)
        p3.font.color.rgb = SLATE_GRAY
        p3.font.name = "Calibri"
        p3.space_before = Pt(3)

    # Prominent Approval & Deferral Branching Box in between rows
    branch_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.3), Inches(11.733), Inches(0.35))
    branch_box.fill.solid()
    branch_box.fill.fore_color.rgb = HIGHLIGHT_BG
    branch_box.line.color.rgb = CARD_BORDER
    branch_box.line.width = Pt(1)
    b_tf = branch_box.text_frame
    b_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    bp = b_tf.paragraphs[0]
    bp.text = "APPROVAL QUEUE BRANCHING:   [APPROVED → Proceed to LOTOTO / SIMOPS Control]   ⇄   [DEFERRED → Clarification Needed → Revise & Resubmit]"
    bp.font.size = Pt(11)
    bp.font.bold = True
    bp.font.color.rgb = AMBER_DEFER
    bp.alignment = PP_ALIGN.CENTER
    bp.font.name = "Calibri"

    add_bottom_takeaway(s4, "Core Principle: The Permit is the central record connecting authorization, safety controls, execution, evidence, and closure.")

    # =========================================================================
    # SLIDE 5: 5. SIMOPS & Safety Incident Protocol
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    setup_slide_header(
        s5, "5", "5. SIMOPS & Safety Incident Protocol",
        "Key Message: Dual operational safeguards monitoring concurrent active work conflicts and capturing safety events in real time."
    )

    sec_w = Inches(5.75)
    sec_h = Inches(4.1)
    sec_top = Inches(1.85)

    # LEFT SECTION: SIMOPS — SIMULTANEOUS OPERATIONS
    left_x = Inches(0.8)
    sim_card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, sec_top, sec_w, sec_h)
    sim_card.fill.solid()
    sim_card.fill.fore_color.rgb = CARD_BG
    sim_card.line.color.rgb = CARD_BORDER
    sim_card.line.width = Pt(1.5)

    sim_hdr = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_x, sec_top, sec_w, Inches(0.55))
    sim_hdr.fill.solid()
    sim_hdr.fill.fore_color.rgb = DEEP_BLUE
    sim_hdr.line.fill.background()
    shtf = sim_hdr.text_frame
    shtf.vertical_anchor = MSO_ANCHOR.MIDDLE
    shp = shtf.paragraphs[0]
    shp.text = "SIMOPS — SIMULTANEOUS OPERATIONS"
    shp.font.size = Pt(13)
    shp.font.bold = True
    shp.font.color.rgb = WHITE
    shp.alignment = PP_ALIGN.CENTER
    shp.font.name = "Calibri"

    stb = s5.shapes.add_textbox(left_x + Inches(0.25), sec_top + Inches(0.65), sec_w - Inches(0.5), sec_h - Inches(0.75))
    stf = stb.text_frame
    stf.word_wrap = True
    stf.margin_left = stf.margin_right = stf.margin_top = stf.margin_bottom = 0

    sp_lead = stf.paragraphs[0]
    sp_lead.text = "Visibility into potential conflicts between concurrent work activities."
    sp_lead.font.size = Pt(12)
    sp_lead.font.bold = True
    sp_lead.font.color.rgb = ACCENT_TEAL
    sp_lead.font.name = "Calibri"
    sp_lead.space_after = Pt(8)

    sim_steps = [
        ("1. Active Permits Scan", "Continuously evaluates all active and approved permits across the facility."),
        ("2. Multi-Vector Comparison", "Checks 4 parameters: Location (Area/Plant), Equipment, Schedule Window, and Permit Type."),
        ("3. Conflict Detection", "Flags simultaneous operations sharing identical zones, machinery, or elevated hazards."),
        ("4. Severity Assessment", "Categorizes detected conflicts by risk tier: High Severity, Medium Severity, or Low Severity."),
        ("5. Live Dashboard & Alert", "Displays open conflicts count, severity matrix, conflict types, and operational review alerts.")
    ]
    for s_title, s_desc in sim_steps:
        p1 = stf.add_paragraph()
        p1.text = f"• {s_title}: {s_desc}"
        p1.font.size = Pt(11)
        p1.font.color.rgb = NAVY
        p1.font.name = "Calibri"
        p1.space_after = Pt(5)

    sp_bv = stf.add_paragraph()
    sp_bv.text = "Value: Helps operational teams identify and mitigate concurrent activity hazards."
    sp_bv.font.size = Pt(11)
    sp_bv.font.bold = True
    sp_bv.font.color.rgb = SLATE_GRAY
    sp_bv.font.name = "Calibri"
    sp_bv.space_before = Pt(4)

    # RIGHT SECTION: SAFETY INCIDENT PROTOCOL
    right_x = Inches(6.783)
    inc_card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_x, sec_top, sec_w, sec_h)
    inc_card.fill.solid()
    inc_card.fill.fore_color.rgb = CARD_BG
    inc_card.line.color.rgb = CARD_BORDER
    inc_card.line.width = Pt(1.5)

    inc_hdr = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, right_x, sec_top, sec_w, Inches(0.55))
    inc_hdr.fill.solid()
    inc_hdr.fill.fore_color.rgb = ACCENT_TEAL
    inc_hdr.line.fill.background()
    ihtf = inc_hdr.text_frame
    ihtf.vertical_anchor = MSO_ANCHOR.MIDDLE
    ihp = ihtf.paragraphs[0]
    ihp.text = "SAFETY INCIDENT PROTOCOL"
    ihp.font.size = Pt(13)
    ihp.font.bold = True
    ihp.font.color.rgb = WHITE
    ihp.alignment = PP_ALIGN.CENTER
    ihp.font.name = "Calibri"

    itb = s5.shapes.add_textbox(right_x + Inches(0.25), sec_top + Inches(0.65), sec_w - Inches(0.5), sec_h - Inches(0.75))
    itf = itb.text_frame
    itf.word_wrap = True
    itf.margin_left = itf.margin_right = itf.margin_top = itf.margin_bottom = 0

    ip_lead = itf.paragraphs[0]
    ip_lead.text = "Structured process for recording and tracking incidents, near misses, and unsafe conditions."
    ip_lead.font.size = Pt(12)
    ip_lead.font.bold = True
    ip_lead.font.color.rgb = ACCENT_TEAL
    ip_lead.font.name = "Calibri"
    ip_lead.space_after = Pt(8)

    inc_steps = [
        ("1. Event Identification", "Captures three distinct categories: Incident, Near Miss, or Unsafe Condition."),
        ("2. Detailed Reporting", "Captures Type, Title, Description, Location, Occurred At timestamp, and optional Linked Permit."),
        ("3. Digital Submission", "Instant submission registers the safety event into the operational tracking registry."),
        ("4. Status Tracking", "Tracks safety lifecycle through transparent stages: Open → Pending HOD Decision → Closed."),
        ("5. Operational Linkage", "Correlates safety events to active work scopes, locations, and organizational units.")
    ]
    for i_title, i_desc in inc_steps:
        p1 = itf.add_paragraph()
        p1.text = f"• {i_title}: {i_desc}"
        p1.font.size = Pt(11)
        p1.font.color.rgb = NAVY
        p1.font.name = "Calibri"
        p1.space_after = Pt(5)

    ip_bv = itf.add_paragraph()
    ip_bv.text = "Value: Enforces transparent safety logging and continuous operational awareness."
    ip_bv.font.size = Pt(11)
    ip_bv.font.bold = True
    ip_bv.font.color.rgb = SLATE_GRAY
    ip_bv.font.name = "Calibri"
    ip_bv.space_before = Pt(4)

    # Bottom Integrated Safeguard Strip
    bridge_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.05), Inches(11.733), Inches(0.45))
    bridge_box.fill.solid()
    bridge_box.fill.fore_color.rgb = WHITE
    bridge_box.line.color.rgb = CARD_BORDER
    bridge_box.line.width = Pt(1)
    br_tf = bridge_box.text_frame
    br_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    br_p = br_tf.paragraphs[0]
    br_p.text = "PTW APPLICATION INTEGRATION:   SIMOPS → Operational Conflict Monitoring   |   INCIDENTS → Safety Event Tracking"
    br_p.font.size = Pt(12)
    br_p.font.bold = True
    br_p.font.color.rgb = NAVY
    br_p.alignment = PP_ALIGN.CENTER
    br_p.font.name = "Calibri"

    add_bottom_takeaway(s5, "Takeaway: Together, SIMOPS and incident management strengthen operational visibility around active work.")

    # Save presentations
    for p in output_paths:
        os.makedirs(os.path.dirname(p), exist_ok=True)
        prs.save(p)
        print(f"Saved: {p}")

if __name__ == "__main__":
    desktop_dir = r"C:\Users\giris\OneDrive\Desktop"
    public_dir = os.path.join(os.getcwd(), "client", "public")
    dist_dir = os.path.join(os.getcwd(), "client", "dist")

    targets = [
        os.path.join(desktop_dir, "PTW_Application_Presentation.pptx"),
        os.path.join(public_dir, "PTW_Application_Presentation.pptx"),
        os.path.join(dist_dir, "PTW_Application_Presentation.pptx"),
        os.path.join(public_dir, "ptw ppt.pptx"),
        os.path.join(dist_dir, "ptw ppt.pptx"),
        os.path.join(public_dir, "Permit_to_Work_PTW_Summary.pptx"),
        os.path.join(dist_dir, "Permit_to_Work_PTW_Summary.pptx")
    ]
    build_ptw_presentation(targets)
