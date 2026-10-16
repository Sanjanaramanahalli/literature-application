import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_clear_simple_ptw_presentation(output_paths):
    prs = Presentation()
    # 16:9 Widescreen: 13.333 x 7.5 inches
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Executive, ultra-clean palette
    NAVY = RGBColor(15, 30, 54)            # Deep dark readable text & headers
    DEEP_BLUE = RGBColor(30, 58, 102)      # Primary badges & headers
    TEAL = RGBColor(13, 148, 136)          # Clean modern accent
    DARK_GRAY = RGBColor(51, 65, 85)       # Easy-reading body text
    MUTED_GRAY = RGBColor(100, 116, 139)   # Sub-labels
    SLATE_GRAY = RGBColor(71, 85, 105)     # Slate header fill
    LIGHT_BG = RGBColor(248, 250, 252)     # Clean light canvas
    WHITE = RGBColor(255, 255, 255)        # Pure card fill
    BORDER_LIGHT = RGBColor(226, 232, 240) # Card outline
    BORDER_ACCENT = RGBColor(13, 148, 136)
    AMBER = RGBColor(217, 119, 6)          # Branching / defer note
    GREEN = RGBColor(16, 149, 106)

    def add_clean_header(slide, title_text, key_idea):
        # Background
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = LIGHT_BG
        bg.line.fill.background()

        # Slim accent line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.04))
        line.fill.solid()
        line.fill.fore_color.rgb = TEAL
        line.line.fill.background()

        # Title block
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.55), Inches(11.733), Inches(1.05))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = title_text
        p1.font.size = Pt(32)
        p1.font.bold = True
        p1.font.color.rgb = NAVY
        p1.font.name = "Calibri"

        p2 = tf.add_paragraph()
        p2.text = f"Key Message: {key_idea}"
        p2.font.size = Pt(14)
        p2.font.color.rgb = MUTED_GRAY
        p2.font.name = "Calibri"
        p2.space_before = Pt(3)

    def add_bottom_takeaway(slide, text):
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.6), Inches(11.733), Inches(0.55))
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = BORDER_LIGHT
        box.line.width = Pt(1.5)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(13.5)
        p.font.bold = True
        p.font.color.rgb = NAVY
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Calibri"

    # =========================================================================
    # SLIDE 1: 1. What is the Application About?
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    add_clean_header(
        s1,
        "1. What is the Application About?",
        "A centralized system to authorize, control, and track work from start to finish."
    )

    # Big 5 Core Pillars Banner
    banner = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(0.65))
    banner.fill.solid()
    banner.fill.fore_color.rgb = NAVY
    banner.line.fill.background()
    b_tf = banner.text_frame
    b_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    b_p = b_tf.paragraphs[0]
    b_p.text = "Work Authorization   +   Safety Control   +   Operational Visibility   +   Accountability   +   Traceability"
    b_p.font.size = Pt(14)
    b_p.font.bold = True
    b_p.font.color.rgb = WHITE
    b_p.alignment = PP_ALIGN.CENTER
    b_p.font.name = "Calibri"

    # 3 Clear Visual Columns: Authorize, Control, Monitor
    c_w = Inches(3.75)
    c_gap = Inches(0.24)
    c_y = Inches(2.65)
    c_h = Inches(3.1)

    cards1 = [
        ("AUTHORIZE", TEAL, [
            "Formal Permit Creation",
            "Structured Approval Queue",
            "Active Work Execution",
            "Completion Evidence Upload",
            "Safe Work Order Closure"
        ]),
        ("CONTROL", DEEP_BLUE, [
            "LOTOTO (Energy Isolation)",
            "Mandatory Safety Controls",
            "Pre-Task Safety Checklists",
            "Work Suspension Capability",
            "Handback & Site Verification"
        ]),
        ("MONITOR", TEAL, [
            "Live SIMOPS Conflict Check",
            "Incident & Near-Miss Tracking",
            "Real-Time Plant Dashboard",
            "Complete Activity Audit Log",
            "Operational Analytics & KPIs"
        ])
    ]

    for i, (head, col, items) in enumerate(cards1):
        left = Inches(0.8) + i * (c_w + c_gap)
        card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, c_y, c_w, c_h)
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = BORDER_LIGHT
        card.line.width = Pt(1.5)

        # Header tag
        tag = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, c_y, c_w, Inches(0.5))
        tag.fill.solid()
        tag.fill.fore_color.rgb = col
        tag.line.fill.background()
        t_tf = tag.text_frame
        t_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        t_p = t_tf.paragraphs[0]
        t_p.text = head
        t_p.font.size = Pt(14)
        t_p.font.bold = True
        t_p.font.color.rgb = WHITE
        t_p.alignment = PP_ALIGN.CENTER
        t_p.font.name = "Calibri"

        # Content list
        tb = s1.shapes.add_textbox(left + Inches(0.25), c_y + Inches(0.65), c_w - Inches(0.5), c_h - Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True
        for itm_idx, itm in enumerate(items):
            p = tf.add_paragraph() if itm_idx > 0 else tf.paragraphs[0]
            p.text = f"✔  {itm}"
            p.font.size = Pt(12)
            p.font.color.rgb = DARK_GRAY
            p.font.name = "Calibri"
            p.space_after = Pt(7)

    # Clear Roles Banner
    roles_box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.9), Inches(11.733), Inches(0.52))
    roles_box.fill.solid()
    roles_box.fill.fore_color.rgb = WHITE
    roles_box.line.color.rgb = BORDER_LIGHT
    roles_box.line.width = Pt(1.5)
    r_tf = roles_box.text_frame
    r_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    r_p = r_tf.paragraphs[0]
    r_p.text = "Clear Roles & Ownership:   Platform Admin   |   Org Admin   |   Job Issuer   |   Safety Officer   |   Operator   |   Viewer"
    r_p.font.size = Pt(12.5)
    r_p.font.bold = True
    r_p.font.color.rgb = DEEP_BLUE
    r_p.alignment = PP_ALIGN.CENTER
    r_p.font.name = "Calibri"

    add_bottom_takeaway(s1, "Takeaway: PTW acts as the central operational record connecting work authorization, safety controls, execution, and closure.")

    # =========================================================================
    # SLIDE 2: 2. Features of the Application
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_clean_header(
        s2,
        "2. Features of the Application",
        "Built-in capabilities organized into 4 clear business functions."
    )

    f_cards = [
        ("1. Work Authorization", DEEP_BLUE, [
            "Permit Creation: Fast drafting",
            "Draft Management: Save & edit",
            "Submission: Required checks",
            "Approval Queue: Multi-role review",
            "Approve / Defer: Give feedback",
            "Resubmission: Quick revisions"
        ]),
        ("2. Safety & Control", TEAL, [
            "LOTOTO Plans: Energy lockout",
            "SIMOPS: Live conflict checks",
            "Safety Controls: Plant rules",
            "Incidents: Event logging",
            "Near Misses: Hazard reporting",
            "Unsafe Conditions: Alert team"
        ]),
        ("3. Execution & Audit", DEEP_BLUE, [
            "Work Execution: Active state",
            "Progress Updates: Live notes",
            "Evidence: Upload completion photos",
            "Activity History: Timestamped log",
            "Suspend Work: Stop on risk",
            "Verify & Close: Sign-off & archive"
        ]),
        ("4. Management & KPIs", TEAL, [
            "Workforce: Users & permissions",
            "Roles: Defined competencies",
            "Organization: Plants & locations",
            "Live Analytics: Active counters",
            "Reports: CSV & operational exports",
            "Billing: Tier & subscription"
        ])
    ]

    card_w2 = Inches(2.78)
    card_gap2 = Inches(0.2)
    card_top2 = Inches(1.85)
    card_h2 = Inches(4.55)

    for i, (title, col, items) in enumerate(f_cards):
        left = Inches(0.8) + i * (card_w2 + card_gap2)
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, card_top2, card_w2, card_h2)
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = BORDER_LIGHT
        card.line.width = Pt(1.5)

        # Card header strip
        hdr = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, card_top2, card_w2, Inches(0.55))
        hdr.fill.solid()
        hdr.fill.fore_color.rgb = col
        hdr.line.fill.background()
        htf = hdr.text_frame
        htf.vertical_anchor = MSO_ANCHOR.MIDDLE
        hp = htf.paragraphs[0]
        hp.text = title
        hp.font.size = Pt(13)
        hp.font.bold = True
        hp.font.color.rgb = WHITE
        hp.alignment = PP_ALIGN.CENTER
        hp.font.name = "Calibri"

        # Item list
        tb = s2.shapes.add_textbox(left + Inches(0.18), card_top2 + Inches(0.7), card_w2 - Inches(0.36), card_h2 - Inches(0.85))
        tf = tb.text_frame
        tf.word_wrap = True
        for itm_idx, itm in enumerate(items):
            p = tf.add_paragraph() if itm_idx > 0 else tf.paragraphs[0]
            p.text = f"• {itm}"
            p.font.size = Pt(11.5)
            p.font.color.rgb = DARK_GRAY
            p.font.name = "Calibri"
            p.space_after = Pt(8)

    add_bottom_takeaway(s2, "One Platform  →  From work authorization to operational reporting and compliance")

    # =========================================================================
    # SLIDE 3: 3. Why is this Application Required?
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_clean_header(
        s3,
        "3. Why is this Application Required?",
        "Replacing fragmented manual habits with a structured digital process."
    )

    # Left: Without Control (Problem)
    left_w = Inches(3.6)
    box_l = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.85), left_w, Inches(4.55))
    box_l.fill.solid()
    box_l.fill.fore_color.rgb = WHITE
    box_l.line.color.rgb = BORDER_LIGHT
    box_l.line.width = Pt(1.5)

    l_hdr = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.85), left_w, Inches(0.55))
    l_hdr.fill.solid()
    l_hdr.fill.fore_color.rgb = SLATE_GRAY
    l_hdr.line.fill.background()
    l_htf = l_hdr.text_frame
    l_htf.vertical_anchor = MSO_ANCHOR.MIDDLE
    l_hp = l_htf.paragraphs[0]
    l_hp.text = "WITHOUT STRUCTURED CONTROL"
    l_hp.font.size = Pt(13)
    l_hp.font.bold = True
    l_hp.font.color.rgb = WHITE
    l_hp.alignment = PP_ALIGN.CENTER
    l_hp.font.name = "Calibri"

    tb_l = s3.shapes.add_textbox(Inches(1.0), Inches(2.55), left_w - Inches(0.4), Inches(3.7))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True

    problems = [
        ("Fragmented Information", "Paper permits, emails, and disconnected spreadsheets."),
        ("Limited Operational Visibility", "Management cannot see active jobs in real time."),
        ("Higher Coordination Risk", "Conflicting work tasks occur without prior warning."),
        ("Weak Accountability", "Hard to verify who approved or checked safety controls.")
    ]
    for idx, (t, d) in enumerate(problems):
        p1 = tf_l.add_paragraph() if idx > 0 else tf_l.paragraphs[0]
        p1.text = f"✖  {t}"
        p1.font.size = Pt(12.5)
        p1.font.bold = True
        p1.font.color.rgb = NAVY
        p1.font.name = "Calibri"

        p2 = tf_l.add_paragraph()
        p2.text = d
        p2.font.size = Pt(11)
        p2.font.color.rgb = DARK_GRAY
        p2.font.name = "Calibri"
        p2.space_after = Pt(10)

    # Right: 6 Business Outcomes (With PTW)
    outcomes_w = Inches(7.8)
    outcomes_left = Inches(4.733)
    
    outcomes = [
        ("1. Controlled Work", "Structured digital workflow before and during execution."),
        ("2. Safer Operations", "Built-in energy isolation (LOTOTO) and SIMOPS conflict checks."),
        ("3. Clear Accountability", "Explicit role-based ownership for approvals and closures."),
        ("4. Operational Visibility", "Single live dashboard for permits, execution, and incidents."),
        ("5. Traceability", "Timestamped progress notes, photo evidence, and closed archives."),
        ("6. Data-Driven Management", "Real-time KPIs and CSV exports for safety review and audits.")
    ]

    card_gw = Inches(3.78)
    card_gh = Inches(1.38)
    gap_x = Inches(0.24)
    gap_y = Inches(0.18)

    for i, (title, desc) in enumerate(outcomes):
        col_idx = i % 2
        row_idx = i // 2
        cx = outcomes_left + col_idx * (card_gw + gap_x)
        cy = Inches(1.85) + row_idx * (card_gh + gap_y)

        c = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, card_gw, card_gh)
        c.fill.solid()
        c.fill.fore_color.rgb = WHITE
        c.line.color.rgb = BORDER_LIGHT
        c.line.width = Pt(1.5)

        # Left teal bar
        bar = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, cy, Inches(0.12), card_gh)
        bar.fill.solid()
        bar.fill.fore_color.rgb = TEAL
        bar.line.fill.background()

        tb = s3.shapes.add_textbox(cx + Inches(0.24), cy + Inches(0.12), card_gw - Inches(0.35), card_gh - Inches(0.24))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(12.5)
        p1.font.bold = True
        p1.font.color.rgb = DEEP_BLUE
        p1.font.name = "Calibri"

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = DARK_GRAY
        p2.font.name = "Calibri"
        p2.space_before = Pt(3)

    add_bottom_takeaway(s3, "Business Outcome: Provides a structured digital process to support safer and more controlled work operations.")

    # =========================================================================
    # SLIDE 4: 4. End-to-End Application Workflow
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_clean_header(
        s4,
        "4. End-to-End Application Workflow",
        "The complete path of a permit from draft creation to final verified closure."
    )

    # Prominent Status Flow
    status_box = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.85), Inches(11.733), Inches(0.48))
    status_box.fill.solid()
    status_box.fill.fore_color.rgb = NAVY
    status_box.line.fill.background()
    s_tf = status_box.text_frame
    s_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    s_p = s_tf.paragraphs[0]
    s_p.text = "OFFICIAL STATUS LIFECYCLE:   DRAFT   →   PENDING APPROVAL   →   APPROVED   →   EXECUTION   →   CLOSED"
    s_p.font.size = Pt(13)
    s_p.font.bold = True
    s_p.font.color.rgb = WHITE
    s_p.alignment = PP_ALIGN.CENTER
    s_p.font.name = "Calibri"

    flow_steps = [
        ("1. CREATE", "Permit Draft", "Issuer creates new permit", TEAL),
        ("2. PLAN", "Scope & Crew", "Location, schedule & team", DEEP_BLUE),
        ("3. SUBMIT", "Review & Check", "Submit to Pending Approval", TEAL),
        ("4. APPROVE", "Approval Queue", "Reviewer Approves or Defers", DEEP_BLUE),
        ("5. CONTROL", "LOTOTO & SIMOPS", "Energy isolation & conflicts", TEAL),
        ("6. EXECUTE", "Work Execution", "Progress notes & evidence", DEEP_BLUE),
        ("7. VERIFY", "Verify & Close", "Site check & completion", TEAL),
        ("8. ARCHIVE", "Audit Record", "Searchable history & KPIs", DEEP_BLUE),
    ]

    card_w4 = Inches(2.78)
    card_h4 = Inches(1.68)
    row1_y = Inches(2.55)
    row2_y = Inches(4.7)

    for i, (stg_num, stg_title, stg_desc, col) in enumerate(flow_steps):
        row = i // 4
        col_idx = i % 4
        left = Inches(0.8) + col_idx * (card_w4 + card_gap2)
        top = row1_y if row == 0 else row2_y

        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, card_w4, card_h4)
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = BORDER_LIGHT
        card.line.width = Pt(1.5)

        # Stage Number Badge
        badge = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.12), top + Inches(0.12), Inches(1.1), Inches(0.32))
        badge.fill.solid()
        badge.fill.fore_color.rgb = col
        badge.line.fill.background()
        btf = badge.text_frame
        btf.vertical_anchor = MSO_ANCHOR.MIDDLE
        bp = btf.paragraphs[0]
        bp.text = stg_num
        bp.font.size = Pt(11)
        bp.font.bold = True
        bp.font.color.rgb = WHITE
        bp.alignment = PP_ALIGN.CENTER
        bp.font.name = "Calibri"

        tb = s4.shapes.add_textbox(left + Inches(0.15), top + Inches(0.52), card_w4 - Inches(0.3), card_h4 - Inches(0.6))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = stg_title
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = NAVY
        p1.font.name = "Calibri"

        p2 = tf.add_paragraph()
        p2.text = stg_desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = DARK_GRAY
        p2.font.name = "Calibri"
        p2.space_before = Pt(3)

    # Clear Deferral Branch Box
    branch = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.3), Inches(11.733), Inches(0.34))
    branch.fill.solid()
    branch.fill.fore_color.rgb = WHITE
    branch.line.color.rgb = BORDER_LIGHT
    branch.line.width = Pt(1)
    br_tf = branch.text_frame
    br_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    br_p = br_tf.paragraphs[0]
    br_p.text = "APPROVAL BRANCHING:   [APPROVED → Proceed to Controls & Execution]   ⇄   [DEFERRED → Add Info → Resubmit]"
    br_p.font.size = Pt(11)
    br_p.font.bold = True
    br_p.font.color.rgb = AMBER
    br_p.alignment = PP_ALIGN.CENTER
    br_p.font.name = "Calibri"

    add_bottom_takeaway(s4, "Core Principle: The Permit is the central record connecting authorization, safety controls, execution, evidence, and closure.")

    # =========================================================================
    # SLIDE 5: 5. SIMOPS & Safety Incident Protocol
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_clean_header(
        s5,
        "5. SIMOPS & Safety Incident Protocol",
        "Dual operational tools to monitor active work conflicts and capture safety events."
    )

    sec_w = Inches(5.75)
    sec_h = Inches(4.1)
    sec_top = Inches(1.85)

    # LEFT: SIMOPS
    left_x = Inches(0.8)
    sim_card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, sec_top, sec_w, sec_h)
    sim_card.fill.solid()
    sim_card.fill.fore_color.rgb = WHITE
    sim_card.line.color.rgb = BORDER_LIGHT
    sim_card.line.width = Pt(1.5)

    sim_hdr = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_x, sec_top, sec_w, Inches(0.55))
    sim_hdr.fill.solid()
    sim_hdr.fill.fore_color.rgb = DEEP_BLUE
    sim_hdr.line.fill.background()
    shtf = sim_hdr.text_frame
    shtf.vertical_anchor = MSO_ANCHOR.MIDDLE
    shp = shtf.paragraphs[0]
    shp.text = "SIMOPS — SIMULTANEOUS OPERATIONS"
    shp.font.size = Pt(13)
    shp.font.bold = True
    shp.font.color.rgb = WHITE
    shp.alignment = PP_ALIGN.CENTER
    shp.font.name = "Calibri"

    stb = s5.shapes.add_textbox(left_x + Inches(0.25), sec_top + Inches(0.68), sec_w - Inches(0.5), sec_h - Inches(0.8))
    stf = stb.text_frame
    stf.word_wrap = True

    sp_lead = stf.paragraphs[0]
    sp_lead.text = "Identifies potential conflicts between simultaneous work activities."
    sp_lead.font.size = Pt(12)
    sp_lead.font.bold = True
    sp_lead.font.color.rgb = TEAL
    sp_lead.font.name = "Calibri"
    sp_lead.space_after = Pt(8)

    sim_points = [
        ("Active Permits Scan", "Continuously checks all approved and active work."),
        ("4 Comparison Factors", "Location, Equipment, Schedule Window, and Permit Type."),
        ("Conflict Detection", "Flags overlapping jobs sharing the same area or gear."),
        ("Severity Assessment", "Ranks conflict risk: High, Medium, or Low severity."),
        ("Dashboard Alerts", "Displays open conflicts, risk levels, and alerts for review.")
    ]
    for st, sd in sim_points:
        p = stf.add_paragraph()
        p.text = f"• {st}: {sd}"
        p.font.size = Pt(11)
        p.font.color.rgb = DARK_GRAY
        p.font.name = "Calibri"
        p.space_after = Pt(6)

    # RIGHT: SAFETY INCIDENT PROTOCOL
    right_x = Inches(6.783)
    inc_card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_x, sec_top, sec_w, sec_h)
    inc_card.fill.solid()
    inc_card.fill.fore_color.rgb = WHITE
    inc_card.line.color.rgb = BORDER_LIGHT
    inc_card.line.width = Pt(1.5)

    inc_hdr = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, right_x, sec_top, sec_w, Inches(0.55))
    inc_hdr.fill.solid()
    inc_hdr.fill.fore_color.rgb = TEAL
    inc_hdr.line.fill.background()
    ihtf = inc_hdr.text_frame
    ihtf.vertical_anchor = MSO_ANCHOR.MIDDLE
    ihp = ihtf.paragraphs[0]
    ihp.text = "SAFETY INCIDENT PROTOCOL"
    ihp.font.size = Pt(13)
    ihp.font.bold = True
    ihp.font.color.rgb = WHITE
    ihp.alignment = PP_ALIGN.CENTER
    ihp.font.name = "Calibri"

    itb = s5.shapes.add_textbox(right_x + Inches(0.25), sec_top + Inches(0.68), sec_w - Inches(0.5), sec_h - Inches(0.8))
    itf = itb.text_frame
    itf.word_wrap = True

    ip_lead = itf.paragraphs[0]
    ip_lead.text = "Structured recording of incidents, near misses, and unsafe conditions."
    ip_lead.font.size = Pt(12)
    ip_lead.font.bold = True
    ip_lead.font.color.rgb = TEAL
    ip_lead.font.name = "Calibri"
    ip_lead.space_after = Pt(8)

    inc_points = [
        ("3 Event Types", "Captures Incidents, Near Misses, and Unsafe Conditions."),
        ("Detailed Report", "Type, Title, Description, Location, Timestamp & Linked Permit."),
        ("Instant Submission", "Enters central register immediately upon completion."),
        ("Status Lifecycle", "Tracks status clearly: Open → Pending HOD Decision → Closed."),
        ("Full Traceability", "Correlates safety events to plant locations and work permits.")
    ]
    for it, idesc in inc_points:
        p = itf.add_paragraph()
        p.text = f"• {it}: {idesc}"
        p.font.size = Pt(11)
        p.font.color.rgb = DARK_GRAY
        p.font.name = "Calibri"
        p.space_after = Pt(6)

    # Bottom Bridge
    b_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.05), Inches(11.733), Inches(0.45))
    b_box.fill.solid()
    b_box.fill.fore_color.rgb = WHITE
    b_box.line.color.rgb = BORDER_LIGHT
    b_box.line.width = Pt(1)
    b_tf = b_box.text_frame
    b_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    bp = b_tf.paragraphs[0]
    bp.text = "PTW APPLICATION INTEGRATION:   SIMOPS → Conflict Monitoring   |   INCIDENTS → Safety Event Tracking"
    bp.font.size = Pt(12)
    bp.font.bold = True
    bp.font.color.rgb = NAVY
    bp.alignment = PP_ALIGN.CENTER
    bp.font.name = "Calibri"

    add_bottom_takeaway(s5, "Takeaway: Together, SIMOPS and incident management strengthen operational visibility around active work.")

    # Save to all target paths
    for p in output_paths:
        os.makedirs(os.path.dirname(p), exist_ok=True)
        prs.save(p)
        print(f"Saved: {p}")

if __name__ == "__main__":
    desktop_dir = r"C:\Users\giris\OneDrive\Desktop"
    public_dir = os.path.join(os.getcwd(), "client", "public")
    dist_dir = os.path.join(os.getcwd(), "client", "dist")

    targets = [
        os.path.join(desktop_dir, "PTW_Application_Presentation.pptx"),
        os.path.join(public_dir, "PTW_Application_Presentation.pptx"),
        os.path.join(dist_dir, "PTW_Application_Presentation.pptx"),
        os.path.join(public_dir, "ptw ppt.pptx"),
        os.path.join(dist_dir, "ptw ppt.pptx"),
        os.path.join(public_dir, "Permit_to_Work_PTW_Summary.pptx"),
        os.path.join(dist_dir, "Permit_to_Work_PTW_Summary.pptx")
    ]
    create_clear_simple_ptw_presentation(targets)
