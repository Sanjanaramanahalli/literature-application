import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

issues_data = [
    (1, "Request Access", 'Request Access cannot be submitted. Warning shows: "We couldn\'t reach the server. Check your connection and send the form again."', "Request Access form should submit successfully without connection errors", "P0", "[New]"),
    (2, "LOTO", "File/evidence cannot be uploaded in the Lockout/Tagout field.", "File/evidence should upload successfully in the LOTO field", "P0", "[Related - PTW-023]"),
    (3, "Permit Templates", "Add Section is not working when creating/editing a permit template.", "Add Section button should create a new section in the template", "P1", "[New]"),
    (4, "Permit Creation", "Machinery does not appear in the dropdown initially. It appears only after refreshing the application.", "Machinery dropdown should load immediately without requiring page refresh", "P1", "[Same - PTW-058]"),
    (5, "Incident Management", 'Chat Admin cannot report an incident. Warning shows: "Validation failed, UUID expired."', "Incidents should be reported successfully without UUID expiration error", "P0", "[New]"),
    (6, "Permit Permissions", "Org Admin can perform multiple permit actions—fill permit details, submit, approve, and start work.", "Role-based permissions should enforce Segregation of Duties", "P0", "[New]"),
    (7, "Notifications", "Permit approval notification appears in Messages, but clicking it does not open the specific permit; it opens/shows all permits.", "Clicking notification should open the specific permit directly", "P1", "[New]"),
    (8, "Permit Creation - LOTO", "Add Lockout/Tagout is not working while creating a permit.", "Add Lockout/Tagout should work and attach isolations to the permit", "P0", "[Related - PTW-021]"),
    (9, "Permit Closure - Evidence", 'After adding evidence during permit closure, "Authentication required" is displayed.', "Evidence should submit without authentication error to allow closure", "P0", "[New]"),
    (10, "Profile", "Profile picture cannot be removed, and the selected profile picture does not display.", "Selected picture should display, and remove option should be available", "P2", "[Same - PTW-036]"),
    (11, "Evidence", "Uploaded evidence cannot be opened/viewed.", "Uploaded evidence files/images should open and view properly", "P1", "[Same - PTW-063]"),
    (12, "HOD Notification", "HOD does not receive a notification when the Job Issuer submits a permit and HOD approval is the next step.", "HOD should receive notification when permit requires approval", "P1", "[Same - PTW-059]"),
    (13, "User & Roles", "Temporary Password - Copy button is not working.", "Temporary password should be copied to clipboard successfully", "P1", "[Same - PTW-005]"),
    (14, "Download", "A Download option is required/available for downloading the details, but its expected behavior should be verified.", "Download option should export permit/incident details properly", "P2", "[New]"),
    (15, "Incident - Root Cause", "Clicking Record Root Cause for an unsafe incident does not open the expected popup/form.", "Record Root Cause button should open the investigation form", "P1", "[New]"),
    (16, "Incident - Corrective Action", "Clicking Corrective Action does not open the expected popup/form.", "Corrective Action button should open the CAPA action form", "P1", "[New]"),
    (17, "Incident - Assign", 'In Manager Incidents, clicking Assign results in "Unable to load page."', "Assign button should open assignment screen/modal without error", "P1", "[New]"),
    (18, "Permit Creation - Back", "There is no Back button while creating a permit.", "Back button should be available on each permit creation step", "P2", "[Same - PTW-027]"),
    (19, "Permit Page - Back", "The Permit page does not have a Back button.", "Back button should be available on the Permit page", "P2", "[Same - PTW-002]"),
    (20, "Permit Creation - Next", "After entering Basic Information, the Next button does not respond on the first click. The user has to click it again.", "Next button should respond and navigate immediately on first click", "P2", "[New]")
]

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_docx(filename):
    doc = docx.Document()

    # Set page layout to Landscape Letter
    section = doc.sections[0]
    section.orientation = docx.enum.section.WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.0)
    section.page_height = Inches(8.5)
    section.left_margin = Inches(0.4)
    section.right_margin = Inches(0.4)
    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.4)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("Software Defect Triage & Severity Classification Report")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(13)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(8)
    run_sub = p_sub.add_run("Module: Permit to Work (PTW), LOTO & Incident Management | Priority: P0, P1, P2")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(9)
    run_sub.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    # Table Creation
    # 21 rows (1 header + 20 data), 6 columns
    table = doc.add_table(rows=21, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    set_table_borders(table, "888888")

    # Column widths (Total = 10.2 inches)
    col_widths = [Inches(0.4), Inches(1.5), Inches(3.4), Inches(3.2), Inches(0.6), Inches(1.1)]

    headers = ["No.", "Module", "Bug / Issue", "Expected Result", "Priority", "Status"]

    # Header Row
    hdr_row = table.rows[0]
    hdr_row._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    for idx, text in enumerate(headers):
        cell = hdr_row.cells[idx]
        cell.width = col_widths[idx]
        set_cell_background(cell, "EAEAEA")
        set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(8.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)

    # Data Rows
    for row_idx, item in enumerate(issues_data):
        row = table.rows[row_idx + 1]
        no_val, mod_val, bug_val, exp_val, pri_val, stat_val = item

        vals = [str(no_val), mod_val, bug_val, exp_val, pri_val, stat_val]

        for col_idx, val in enumerate(vals):
            cell = row.cells[col_idx]
            cell.width = col_widths[col_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            set_cell_margins(cell, top=40, bottom=40, left=70, right=70)
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05

            run = p.add_run(val)
            run.font.name = "Arial"
            run.font.size = Pt(8)

            if col_idx == 0:  # No.
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif col_idx == 4:  # Priority
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run.font.bold = True
                if val == "P0":
                    run.font.color.rgb = RGBColor(0x9B, 0x2C, 0x2C)
                elif val == "P1":
                    run.font.color.rgb = RGBColor(0xC0, 0x56, 0x21)
                else:
                    run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
            elif col_idx == 5:  # Status
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)

    doc.save(filename)
    print(f"Generated Word Document: {filename}")

if __name__ == "__main__":
    target = os.path.abspath("Issue_Prioritization_Table.docx")
    build_docx(target)
