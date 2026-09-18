import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_ptw_presentation(output_path):
    prs = Presentation()
    # 16:9 Widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # Theme Colors - subtle, corporate, executive
    NAVY = RGBColor(16, 42, 77)
    AMBER = RGBColor(217, 119, 6)
    DARK_GRAY = RGBColor(51, 65, 85)
    LIGHT_BG = RGBColor(248, 250, 252)
    CARD_BG = RGBColor(255, 255, 255)
    CARD_BORDER = RGBColor(226, 232, 240)
    GREEN = RGBColor(16, 149, 106)
    RED = RGBColor(220, 38, 38)
    BLUE = RGBColor(2, 132, 199)

    def add_header(slide, slide_num, question, subtitle):
        # Background
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = LIGHT_BG
        bg.line.fill.background()

        # Top Accent Line
        top_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.06))
        top_line.fill.solid()
        top_line.fill.fore_color.rgb = AMBER
        top_line.line.fill.background()

        # Question Header Box
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.55), Inches(11.733), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = f"QUESTION {slide_num}: {question.upper()}"
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = AMBER
        p1.font.name = "Arial"

        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(22)
        p2.font.bold = True
        p2.font.color.rgb = NAVY
        p2.font.name = "Arial"

    # ==========================================
    # SLIDE 1: What is the Application All About?
    # ==========================================
    s1 = prs.slides.add_slide(blank_slide_layout)
    add_header(s1, "1", "What is the Application All About?", "Permit to Work (PTW) System — High-Hazard Authorization Platform")

    # Formula Card
    f_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(1.0))
    f_card.fill.solid()
    f_card.fill.fore_color.rgb = NAVY
    f_card.line.fill.background()
    f_tf = f_card.text_frame
    f_tf.word_wrap = True
    f_p = f_tf.paragraphs[0]
    f_p.alignment = PP_ALIGN.CENTER
    f_p.text = "THE CORE SAFETY EQUATION\nPTW = Formal Permission  +  Hazard Identification  +  Safety Controls  +  Authorized Sign-Off"
    f_p.font.size = Pt(15)
    f_p.font.bold = True
    f_p.font.color.rgb = RGBColor(255, 255, 255)

    # 4 Roles Cards
    roles = [
        ("👷 Permit Receiver", "The Job Leader", "Supervisor/Foreman who applies for the permit, manages crew safety, and remains physically present at the site until job completion.", BLUE),
        ("📝 Permit Issuer", "The Plant Owner", "Operations/Area Engineer who inspects the area, isolates energy/piping (LOTO), verifies safety, and signs authorization.", AMBER),
        ("📋 PTW Coordinator", "Traffic Controller", "Central master controller who maintains the Permit Register, tracks expiry, and ensures no two dangerous jobs clash (SIMOPS).", NAVY),
        ("🦺 Safety Officer (EHS)", "Safety Guardian", "Conducts gas testing (LEL/toxic), inspects safety gear, endorses permits, and exercises instant Stop Work Authority.", GREEN),
    ]

    card_w = Inches(2.78)
    card_gap = Inches(0.2)
    card_top = Inches(3.0)
    card_h = Inches(3.9)

    for i, (title, role_type, desc, accent_col) in enumerate(roles):
        left = Inches(0.8) + i * (card_w + card_gap)
        c = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, card_top, card_w, card_h)
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = CARD_BORDER
        c.line.width = Pt(1.5)

        tb = s1.shapes.add_textbox(left + Inches(0.2), card_top + Inches(0.2), card_w - Inches(0.4), card_h - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = accent_col

        p_sub = tf.add_paragraph()
        p_sub.text = f"({role_type})"
        p_sub.font.size = Pt(12)
        p_sub.font.italic = True
        p_sub.font.color.rgb = DARK_GRAY
        p_sub.space_after = Pt(14)

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(13)
        p_desc.font.color.rgb = DARK_GRAY

    # ==========================================
    # SLIDE 2: Features of this Application
    # ==========================================
    s2 = prs.slides.add_slide(blank_slide_layout)
    add_header(s2, "2", "What are the Features of this Application?", "Color-Coded Permits, Real-Time Controls & Workforce Tracking")

    # Left Column: Permit Types
    permits = [
        ("🔴 Hot Work (Red)", "For sparks, fire, welding, cutting, grinding.\n• Must test gas (0.0% LEL), assign fire watch & cover sewers within 75 ft."),
        ("🔵 Cold Work (Light Blue)", "Base permit for work without sparks or flames.\n• Pipe blinding/de-blinding, formwork, painting, cable pulling."),
        ("🟢 Confined Space Entry (Green)", "Tanks, vessels, manholes, silos, boilers.\n• Tests oxygen, toxic gases; mandates blowers & outside standby guard."),
        ("🟡 Electrical & LOTO (Yellow/Orange)", "Work on energized or isolated electrical circuits.\n• Lockout/Tagout padlocks, zero-energy switch checks, circuit grounding."),
        ("🟣 Specialized Permits (Purple/Brown/Black)", "• Height (>1.8m harness), Excavation (underground scan), Night Work."),
    ]

    p_top = Inches(1.8)
    p_h = Inches(5.1)
    p_w = Inches(6.8)

    box_l = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), p_top, p_w, p_h)
    box_l.fill.solid()
    box_l.fill.fore_color.rgb = CARD_BG
    box_l.line.color.rgb = CARD_BORDER
    box_l.line.width = Pt(1.5)

    tb_l = s2.shapes.add_textbox(Inches(1.0), p_top + Inches(0.2), p_w - Inches(0.4), p_h - Inches(0.4))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True

    p_header = tf_l.paragraphs[0]
    p_header.text = "1. Standard Color-Coded Work Permits"
    p_header.font.size = Pt(16)
    p_header.font.bold = True
    p_header.font.color.rgb = NAVY
    p_header.space_after = Pt(10)

    for name, details in permits:
        p_item = tf_l.add_paragraph()
        p_item.text = name
        p_item.font.size = Pt(13)
        p_item.font.bold = True
        p_item.font.color.rgb = NAVY

        p_sub = tf_l.add_paragraph()
        p_sub.text = details
        p_sub.font.size = Pt(11)
        p_sub.font.color.rgb = DARK_GRAY
        p_sub.space_after = Pt(6)

    # Right Column: Core Smart Features
    r_left = Inches(7.8)
    r_w = Inches(4.733)
    box_r = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, r_left, p_top, r_w, p_h)
    box_r.fill.solid()
    box_r.fill.fore_color.rgb = CARD_BG
    box_r.line.color.rgb = CARD_BORDER
    box_r.line.width = Pt(1.5)

    tb_r = s2.shapes.add_textbox(r_left + Inches(0.25), p_top + Inches(0.2), r_w - Inches(0.5), p_h - Inches(0.4))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    p_rhead = tf_r.paragraphs[0]
    p_rhead.text = "2. Built-in Safety Governance Features"
    p_rhead.font.size = Pt(16)
    p_rhead.font.bold = True
    p_rhead.font.color.rgb = NAVY
    p_rhead.space_after = Pt(14)

    smart_features = [
        ("📋 4-Category Hazard Profiling", "Systematic screening of Conditions (height/space), Exposure (noise/heat), Chemicals (COSHH), and Events (explosion/collapse)."),
        ("🗣️ Pre-Task Safety Talks (TBT)", "Compulsory safety briefings held at the worksite and signed by all workers before work begins."),
        ("👥 Crew Acknowledgment Sheet", "Digital registry capturing every worker's name, employee ID, and trade role for legal accountability."),
        ("🔄 Daily Revalidation Engine", "Multi-day permits (up to 6 days) strictly require daily re-inspection and Issuer signature each morning."),
    ]

    for sf_title, sf_desc in smart_features:
        p_t = tf_r.add_paragraph()
        p_t.text = sf_title
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = AMBER

        p_d = tf_r.add_paragraph()
        p_d.text = sf_desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = DARK_GRAY
        p_d.space_after = Pt(10)

    # ==========================================
    # SLIDE 3: Why is This Application Required?
    # ==========================================
    s3 = prs.slides.add_slide(blank_slide_layout)
    add_header(s3, "3", "Why is This Application Required?", "Catastrophe Prevention, Energy Isolation & Regulatory Compliance")

    reasons = [
        ("1. Disaster & Explosion Prevention", "Refineries hold explosive gases and volatile hydrocarbons. A single stray spark near an unmonitored line causes catastrophic fires. PTW proves zero flammable gas exists before tools touch equipment.", RED),
        ("2. Conflict Detection (No Blind Spots)", "Multiple contractor crews work on the same plant simultaneously. PTW deconflicts overlapping zones (e.g. stopping hot welding directly over toxic gas venting).", BLUE),
        ("3. Energy Isolation (LOTO Verification)", "Ensures equipment is physically de-energized, depressurized, drained, and locked out so machines cannot start by accident while workers are inside.", AMBER),
        ("4. Shared Multi-Party Responsibility", "Worksite safety is never placed on a single person. Issuer, Receiver, and Safety Officers must all jointly inspect and sign off before work starts.", NAVY),
        ("5. Regulatory Compliance & Traceability", "Maintains an unalterable audit log of gas readings, safety talks, and approvals to satisfy OSHA, ISO 45001, and corporate safety regulations.", GREEN),
    ]

    r_top = Inches(1.8)
    r_card_h = Inches(0.92)
    r_gap = Inches(0.12)

    for i, (title, desc, accent) in enumerate(reasons):
        c_top = r_top + i * (r_card_h + r_gap)
        c = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), c_top, Inches(11.733), r_card_h)
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = CARD_BORDER
        c.line.width = Pt(1.5)

        # Left color strip
        strip = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), c_top, Inches(0.18), r_card_h)
        strip.fill.solid()
        strip.fill.fore_color.rgb = accent
        strip.line.fill.background()

        tb = s3.shapes.add_textbox(Inches(1.2), c_top + Inches(0.1), Inches(11.1), r_card_h - Inches(0.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = NAVY

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = DARK_GRAY

    # ==========================================
    # SLIDE 4: How Does the Application Flow End-to-End?
    # ==========================================
    s4 = prs.slides.add_slide(blank_slide_layout)
    add_header(s4, "4", "How Does the Application Flow End-to-End?", "The 8-Stage Lifecycle from Job Planning to Final Site Handback")

    steps = [
        ("Step 1", "Permit Request", "Receiver submits task scope, tools, JSA, and crew list.", BLUE),
        ("Step 2", "Technical Review", "Issuer verifies equipment isolations, valves, and switches.", NAVY),
        ("Step 3", "SIMOPS Check", "Coordinator registers permit and checks nearby work zones.", AMBER),
        ("Step 4", "Safety Inspection", "Safety Officer tests air quality (gas test) & endorses.", GREEN),
        ("Step 5", "Activation & TBT", "Permit placed in site box; crew briefed in Toolbox Talk.", BLUE),
        ("Step 6", "Work & Revalidation", "Job executed; daily review and signature for multi-day jobs.", NAVY),
        ("Step 7", "Site Handback", "Tools packed, isolations removed, site cleaned thoroughly.", AMBER),
        ("Step 8", "Final Closeout", "Issuer & Receiver sign completion; permit archived.", GREEN),
    ]

    step_w = Inches(2.78)
    step_h = Inches(2.35)
    row1_top = Inches(1.85)
    row2_top = Inches(4.45)

    for i, (st, sname, sdesc, col) in enumerate(steps):
        row = i // 4
        col_idx = i % 4
        left = Inches(0.8) + col_idx * (step_w + card_gap)
        top = row1_top if row == 0 else row2_top

        c = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, step_w, step_h)
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = CARD_BORDER
        c.line.width = Pt(1.5)

        # Header tag inside card
        tag = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.15), top + Inches(0.15), Inches(0.9), Inches(0.32))
        tag.fill.solid()
        tag.fill.fore_color.rgb = col
        tag.line.fill.background()
        tag.text_frame.word_wrap = True
        tag_p = tag.text_frame.paragraphs[0]
        tag_p.text = st
        tag_p.font.size = Pt(11)
        tag_p.font.bold = True
        tag_p.font.color.rgb = RGBColor(255, 255, 255)
        tag_p.alignment = PP_ALIGN.CENTER

        tb = s4.shapes.add_textbox(left + Inches(0.15), top + Inches(0.55), step_w - Inches(0.3), step_h - Inches(0.65))
        tf = tb.text_frame
        tf.word_wrap = True

        p_name = tf.paragraphs[0]
        p_name.text = sname
        p_name.font.size = Pt(14)
        p_name.font.bold = True
        p_name.font.color.rgb = NAVY
        p_name.space_after = Pt(6)

        p_desc = tf.add_paragraph()
        p_desc.text = sdesc
        p_desc.font.size = Pt(11.5)
        p_desc.font.color.rgb = DARK_GRAY

    # ==========================================
    # SLIDE 5: What Happens During Simultaneous Incidents?
    # ==========================================
    s5 = prs.slides.add_slide(blank_slide_layout)
    add_header(s5, "5", "What Happens During Simultaneous Incidents?", "SIMOPS Conflict Deconfliction & Stop Work Authority (SWA) Protocols")

    # Left Box: SIMOPS
    s5_top = Inches(1.8)
    s5_h = Inches(5.1)
    s5_w = Inches(5.6)

    box_s5_l = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), s5_top, s5_w, s5_h)
    box_s5_l.fill.solid()
    box_s5_l.fill.fore_color.rgb = CARD_BG
    box_s5_l.line.color.rgb = CARD_BORDER
    box_s5_l.line.width = Pt(1.5)

    tb_s5_l = s5.shapes.add_textbox(Inches(1.0), s5_top + Inches(0.2), s5_w - Inches(0.4), s5_h - Inches(0.4))
    tf_s5_l = tb_s5_l.text_frame
    tf_s5_l.word_wrap = True

    p_l1 = tf_s5_l.paragraphs[0]
    p_l1.text = "1. SIMOPS (Simultaneous Operations)"
    p_l1.font.size = Pt(16)
    p_l1.font.bold = True
    p_l1.font.color.rgb = NAVY
    p_l1.space_after = Pt(8)

    simops_items = [
        ("What is SIMOPS?", "When two or more teams work at the same time in the same or nearby area (e.g. welding on scaffolding directly above workers on a pipeline)."),
        ("Automatic Conflict Lock", "The Coordinator checks the master map. If jobs clash (Vertical elevation hazard or Horizontal proximity), the system locks permit activation."),
        ("Mandatory SIMOPS Checklist", "Crews must attach a joint checklist establishing physical fire blankets, protective netting, dedicated watchmen, and shared radio channels."),
        ("Joint Sign-Off Required", "Both work crew leaders must agree and co-sign before either team can begin."),
    ]

    for st, sd in simops_items:
        p_t = tf_s5_l.add_paragraph()
        p_t.text = f"• {st}"
        p_t.font.size = Pt(12.5)
        p_t.font.bold = True
        p_t.font.color.rgb = AMBER

        p_d = tf_s5_l.add_paragraph()
        p_d.text = sd
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = DARK_GRAY
        p_d.space_after = Pt(6)

    # Right Box: Safety Incidents / SWA
    s5_r_left = Inches(6.8)
    s5_r_w = Inches(5.733)

    box_s5_r = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, s5_r_left, s5_top, s5_r_w, s5_h)
    box_s5_r.fill.solid()
    box_s5_r.fill.fore_color.rgb = CARD_BG
    box_s5_r.line.color.rgb = CARD_BORDER
    box_s5_r.line.width = Pt(1.5)

    tb_s5_r = s5.shapes.add_textbox(s5_r_left + Inches(0.2), s5_top + Inches(0.2), s5_r_w - Inches(0.4), s5_h - Inches(0.4))
    tf_s5_r = tb_s5_r.text_frame
    tf_s5_r.word_wrap = True

    p_r1 = tf_s5_r.paragraphs[0]
    p_r1.text = "2. Emergency / Incident Response Protocol"
    p_r1.font.size = Pt(16)
    p_r1.font.bold = True
    p_r1.font.color.rgb = RED
    p_r1.space_after = Pt(8)

    incident_steps = [
        ("🛑 Step 1: Stop Work Authority (SWA)", "ANY worker can stop work immediately without asking permission if gas leaks, bad weather, chemical spills, or alarms occur."),
        ("🚨 Step 2: Notify & Sound Alarm", "Permit Receiver alerts the Control Room, Issuer, and Safety team. Crew evacuates to safe muster points if needed."),
        ("❌ Step 3: All Active Permits Voided", "When an emergency alarm rings, ALL active permits in the area are instantly cancelled and declared invalid. No work can restart informally."),
        ("🔍 Step 4: Investigate & Closeout", "Incident cause is investigated (Root Cause Analysis). Existing permit is officially stamped 'Work Suspended'."),
        ("🔄 Step 5: New Permit Required", "Work can ONLY resume by issuing a BRAND NEW PERMIT with fresh site inspections and gas tests."),
    ]

    for it, idesc in incident_steps:
        p_t = tf_s5_r.add_paragraph()
        p_t.text = it
        p_t.font.size = Pt(12)
        p_t.font.bold = True
        p_t.font.color.rgb = NAVY

        p_d = tf_s5_r.add_paragraph()
        p_d.text = idesc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = DARK_GRAY
        p_d.space_after = Pt(6)

    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    out_dir = os.path.join(os.getcwd(), "client", "public")
    os.makedirs(out_dir, exist_ok=True)
    out_file = os.path.join(out_dir, "Permit_to_Work_PTW_Summary.pptx")
    create_ptw_presentation(out_file)
